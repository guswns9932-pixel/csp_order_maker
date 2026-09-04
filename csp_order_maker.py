# -*- coding: utf-8 -*-
"""
CSP 주문접수 업로드 파일 생성기  (알파 v0.1)

사용법
  1) 이 파일과 'CSP_주문접수_업로드_통합양식.xlsx' 를 같은 폴더에 둔다
  2) python csp_order_maker.py
  3) 공통값을 채우고 -> 품목 라인을 추가 -> [엑셀 파일 생성]

필요 패키지 : openpyxl
"""

import os
import sys
import json
import datetime as dt
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment
from openpyxl.utils import get_column_letter

APP_TITLE = "CSP 주문접수 업로드 파일 생성기  (alpha v0.1)"
TEMPLATE_NAME = "CSP_주문접수_업로드_통합양식.xlsx"
SETTINGS_NAME = "csp_order_maker_settings.json"

# ---------------------------------------------------------------- 양식 정의
# (엑셀열, 헤더명, 필수여부)
COLUMNS = [
    ("A", "판매오더유형", True),
    ("B", "판매처코드", True),
    ("C", "인도처코드", True),
    ("D", "유통경로", True),
    ("E", "제품군", False),
    ("F", "고객PO번호", True),
    ("G", "고객PO일자", False),
    ("H", "인도조건", True),
    ("I", "인도장소", False),
    ("J", "가격결정일", True),
    ("K", "통화", True),
    ("L", "고객라인", False),
    ("M", "고객설비", False),
    ("N", "고객공정", False),
    ("O", "고객세부공정", False),
    ("P", "고객설비호기", True),
    ("Q", "자재코드", True),
    ("R", "오더수량", True),
    ("S", "단위", False),
    ("T", "납품요청일", True),
    ("U", "출하지점", False),
    ("V", "조건유형", False),
    ("W", "단가", False),
    ("X", "금액", True),
    ("Y", "통신유형", True),
]

# 모든 행이 같은 값을 갖는 항목 (화면 위쪽에서 한 번만 입력)
COMMON_KEYS = ["A", "B", "C", "D", "E", "G", "H", "I", "J", "K",
               "L", "M", "R", "S", "U", "V", "Y"]
# 행마다 달라지는 항목 (아래 표에서 행별 입력)
LINE_KEYS = ["F", "N", "O", "P", "Q", "T", "W", "X"]

HEADER_BY_KEY = {k: h for k, h, _ in COLUMNS}
REQUIRED_KEYS = {k for k, _, r in COLUMNS if r}


def app_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


# ---------------------------------------------------------------- 마스터 데이터
class MasterData:
    """통합양식 파일의 코드 시트 / FSC 시트를 읽어들인다."""

    def __init__(self, path):
        self.path = path
        self.order_types = []      # [(코드, 내역)]
        self.sold_to = []          # [(코드, 명, 주소)]
        self.ship_to = []
        self.channels = []
        self.inco_terms = []
        self.currencies = []
        self.comm_types = []
        self.fsc = []              # [(FSC, VER, FSC NM, 설명, 상태)]
        self._load()

    @staticmethod
    def _s(v):
        return "" if v is None else str(v).strip()

    def _code_sheet(self, wb, name, header_rows=2):
        """A=코드, B=내역 형태의 시트를 읽는다."""
        out = []
        if name not in wb.sheetnames:
            return out
        ws = wb[name]
        for row in ws.iter_rows(min_row=header_rows + 1, max_col=2, values_only=True):
            code = self._s(row[0])
            if not code:
                continue
            out.append((code, self._s(row[1] if len(row) > 1 else "")))
        return out

    def _partner_sheet(self, wb, name):
        out = []
        if name not in wb.sheetnames:
            return out
        ws = wb[name]
        for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
            code = self._s(row[0])
            if not code:
                continue
            out.append((code, self._s(row[1]), self._s(row[2])))
        return out

    def _load(self):
        wb = load_workbook(self.path, read_only=True, data_only=True)
        try:
            self.order_types = self._code_sheet(wb, "판매오더유형")
            self.channels = self._code_sheet(wb, "유통경로")
            self.inco_terms = self._code_sheet(wb, "인도조건")
            self.currencies = self._code_sheet(wb, "통화")
            self.comm_types = self._code_sheet(wb, "통신유형")
            self.sold_to = self._partner_sheet(wb, "판매처코드")
            self.ship_to = self._partner_sheet(wb, "인도처코드")

            if "FSC" in wb.sheetnames:
                ws = wb["FSC"]
                seen = set()
                for row in ws.iter_rows(min_row=2, max_col=11, values_only=True):
                    code = self._s(row[1])          # B열 : FSC
                    if not code or code in seen:
                        continue
                    seen.add(code)
                    self.fsc.append((
                        code,
                        self._s(row[2]),            # C열 : VER
                        self._s(row[5]),            # F열 : FSC NM
                        self._s(row[7]).replace("\n", " "),   # H열 : 설명
                        self._s(row[10]),           # K열 : 상태
                    ))
        finally:
            wb.close()

    @property
    def fsc_codes(self):
        return {f[0] for f in self.fsc}


# ---------------------------------------------------------------- 값 변환 / 검증
def parse_date(text):
    """YYYYMMDD / YYYY-MM-DD / YYYY.MM.DD / YYYY/MM/DD -> datetime.date"""
    t = str(text).strip().replace(".", "-").replace("/", "-")
    if not t:
        return None
    for fmt in ("%Y-%m-%d", "%Y%m%d"):
        try:
            return dt.datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    return None


def parse_int(text):
    t = str(text).strip().replace(",", "")
    if not t:
        return None
    try:
        return int(float(t))
    except ValueError:
        return None


def combo_code(text):
    """'ZOR1 - 제품 일반주문' 형태에서 코드만 뽑아낸다."""
    return str(text).split(" - ", 1)[0].strip()


# ---------------------------------------------------------------- 엑셀 출력
def build_output(template_path, rows, out_path):
    """rows : [{열키: 값}] 을 받아 Sheet1 양식의 새 파일을 만든다."""
    tpl = load_workbook(template_path)
    tws = tpl["Sheet1"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # --- 헤더행 : 원본 서식 그대로 복사
    for idx, (key, header, _) in enumerate(COLUMNS, start=1):
        src = tws.cell(row=1, column=idx)
        dst = ws.cell(row=1, column=idx, value=src.value)
        dst.font = Font(name=src.font.name, sz=src.font.sz, b=src.font.b,
                        color=src.font.color)
        if src.fill and src.fill.fill_type:
            dst.fill = PatternFill(fill_type=src.fill.fill_type,
                                   fgColor=src.fill.fgColor,
                                   bgColor=src.fill.bgColor)
        dst.border = Border(left=src.border.left, right=src.border.right,
                            top=src.border.top, bottom=src.border.bottom)
        dst.alignment = Alignment(horizontal=src.alignment.horizontal,
                                  vertical=src.alignment.vertical,
                                  wrap_text=src.alignment.wrap_text)
        letter = get_column_letter(idx)
        if tws.column_dimensions[letter].width:
            ws.column_dimensions[letter].width = tws.column_dimensions[letter].width

    # --- 데이터행
    for r, data in enumerate(rows, start=2):
        for idx, (key, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=idx, value=data.get(key))
            if key in ("G", "J"):          # 텍스트 형식 YYYYMMDD
                cell.number_format = "@"
            elif key == "T":               # 엑셀 날짜값
                cell.number_format = "yyyy-mm-dd"

    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------- 검색 팝업
class PickerDialog(tk.Toplevel):
    """검색 + 목록 선택 공용 팝업."""

    def __init__(self, parent, title, columns, widths, rows, key_index=0):
        super().__init__(parent)
        self.title(title)
        self.transient(parent)
        self.grab_set()
        self.result = None
        self._rows = rows
        self._key_index = key_index

        top = ttk.Frame(self, padding=8)
        top.pack(fill="x")
        ttk.Label(top, text="검색").pack(side="left")
        self.var = tk.StringVar()
        ent = ttk.Entry(top, textvariable=self.var, width=40)
        ent.pack(side="left", padx=6)
        ent.focus_set()
        self.var.trace_add("write", lambda *_: self._refresh())
        self.count = ttk.Label(top, text="")
        self.count.pack(side="left", padx=6)

        body = ttk.Frame(self, padding=(8, 0, 8, 8))
        body.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(body, columns=columns, show="headings",
                                 height=18, selectmode="browse")
        for c, w in zip(columns, widths):
            self.tree.heading(c, text=c)
            self.tree.column(c, width=w, anchor="w")
        vs = ttk.Scrollbar(body, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vs.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vs.pack(side="left", fill="y")
        self.tree.bind("<Double-1>", lambda e: self._ok())
        self.tree.bind("<Return>", lambda e: self._ok())
        ent.bind("<Return>", lambda e: self._focus_first())
        ent.bind("<Down>", lambda e: self._focus_first())

        btn = ttk.Frame(self, padding=(8, 0, 8, 8))
        btn.pack(fill="x")
        ttk.Button(btn, text="선택", command=self._ok).pack(side="right")
        ttk.Button(btn, text="취소", command=self.destroy).pack(side="right", padx=6)

        self._refresh()
        self.geometry("+%d+%d" % (parent.winfo_rootx() + 60, parent.winfo_rooty() + 60))

    def _refresh(self):
        kw = self.var.get().strip().lower()
        self.tree.delete(*self.tree.get_children())
        shown = 0
        for row in self._rows:
            if kw and not any(kw in str(v).lower() for v in row):
                continue
            self.tree.insert("", "end", values=row)
            shown += 1
            if shown >= 500:
                break
        self.count.config(text="%d건 표시 (전체 %d건)" % (shown, len(self._rows)))

    def _focus_first(self):
        kids = self.tree.get_children()
        if kids:
            self.tree.selection_set(kids[0])
            self.tree.focus(kids[0])
            self.tree.focus_set()

    def _ok(self):
        sel = self.tree.selection()
        if not sel:
            return
        self.result = self.tree.item(sel[0], "values")[self._key_index]
        self.destroy()


# ---------------------------------------------------------------- 메인 앱
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1280x820")
        self.minsize(1100, 700)

        try:
            self.option_add("*Font", ("맑은 고딕", 9))
        except tk.TclError:
            pass

        self.master_path = tk.StringVar(value=self._find_template())
        self.md = None
        self.common_vars = {}
        self.line_vars = {}
        self.lines = []          # [{열키: 원시 문자열}]
        self.editing_index = None

        self._build_ui()
        self._load_master(initial=True)
        self._load_settings()

    # ---------- 초기화
    def _find_template(self):
        p = os.path.join(app_dir(), TEMPLATE_NAME)
        return p if os.path.exists(p) else ""

    def _load_master(self, initial=False):
        path = self.master_path.get()
        if not path or not os.path.exists(path):
            if not initial:
                messagebox.showerror("오류", "양식 파일을 찾을 수 없습니다.")
            self.status.config(text="양식 파일을 지정해 주세요.")
            return
        try:
            self.md = MasterData(path)
        except Exception as e:
            messagebox.showerror("오류", "양식 파일을 읽지 못했습니다.\n\n%s" % e)
            return
        self._fill_combos()
        self.status.config(
            text="양식 로드 완료 · 판매처 %d · 인도처 %d · FSC %d건"
                 % (len(self.md.sold_to), len(self.md.ship_to), len(self.md.fsc)))

    def _fill_combos(self):
        def items(pairs):
            return ["%s - %s" % (c, d) if d else c for c, d in pairs]

        self.cbo["A"]["values"] = items(self.md.order_types)
        self.cbo["D"]["values"] = items(self.md.channels)
        self.cbo["H"]["values"] = items(self.md.inco_terms)
        self.cbo["K"]["values"] = items(self.md.currencies)
        self.cbo["Y"]["values"] = items(self.md.comm_types)

    # ---------- 화면 구성
    def _build_ui(self):
        root = ttk.Frame(self, padding=8)
        root.pack(fill="both", expand=True)

        # 양식 파일
        bar = ttk.Frame(root)
        bar.pack(fill="x", pady=(0, 6))
        ttk.Label(bar, text="양식 파일").pack(side="left")
        ttk.Entry(bar, textvariable=self.master_path).pack(
            side="left", fill="x", expand=True, padx=6)
        ttk.Button(bar, text="찾아보기", command=self._pick_master).pack(side="left")
        ttk.Button(bar, text="다시 읽기",
                   command=lambda: self._load_master()).pack(side="left", padx=4)

        # 공통값
        box = ttk.LabelFrame(root, text=" 공통값 (모든 행에 동일하게 들어감) ",
                             padding=8)
        box.pack(fill="x")
        self.cbo = {}
        self._common_grid(box)

        # 품목 라인 입력
        lbox = ttk.LabelFrame(root, text=" 품목 라인 (행마다 달라지는 값) ", padding=8)
        lbox.pack(fill="both", expand=True, pady=(8, 0))
        self._line_form(lbox)
        self._line_table(lbox)

        # 하단
        bottom = ttk.Frame(root)
        bottom.pack(fill="x", pady=(8, 0))
        self.status = ttk.Label(bottom, text="", foreground="#555")
        self.status.pack(side="left")
        ttk.Button(bottom, text="엑셀 파일 생성",
                   command=self._export).pack(side="right")
        ttk.Button(bottom, text="공통값 저장",
                   command=self._save_settings).pack(side="right", padx=6)

    def _common_grid(self, parent):
        """공통값 입력칸을 4열로 배치."""
        specs = [
            ("A", "combo"), ("B", "pick_sold"), ("C", "pick_ship"), ("D", "combo"),
            ("E", "entry"), ("G", "date8"), ("H", "combo"), ("I", "entry"),
            ("J", "date8"), ("K", "combo"), ("L", "entry"), ("M", "entry"),
            ("R", "fixed"), ("S", "entry"), ("U", "entry"), ("V", "entry"),
            ("Y", "combo"),
        ]
        defaults = {"E": "10", "R": "1", "S": "EA", "U": "1100", "V": "PR00"}
        for i, (key, kind) in enumerate(specs):
            r, c = divmod(i, 4)
            cell = ttk.Frame(parent)
            cell.grid(row=r, column=c, sticky="ew", padx=6, pady=3)
            parent.columnconfigure(c, weight=1)

            label = HEADER_BY_KEY[key]
            if key in REQUIRED_KEYS:
                label = "* " + label
            ttk.Label(cell, text="%s (%s)" % (label, key), width=16).pack(side="left")

            var = tk.StringVar(value=defaults.get(key, ""))
            self.common_vars[key] = var

            if kind == "combo":
                w = ttk.Combobox(cell, textvariable=var, state="readonly", width=22)
                w.pack(side="left", fill="x", expand=True)
                self.cbo[key] = w
            elif kind in ("pick_sold", "pick_ship"):
                ttk.Entry(cell, textvariable=var, width=12).pack(side="left")
                lbl = ttk.Label(cell, text="", foreground="#0a6", width=14)
                lbl.pack(side="left", padx=3)
                setattr(self, "_name_%s" % key, lbl)
                which = "sold" if kind == "pick_sold" else "ship"
                ttk.Button(cell, text="찾기", width=5,
                           command=lambda w=which: self._pick_partner(w)).pack(side="left")
                var.trace_add("write", lambda *_a, k=key: self._show_partner_name(k))
            elif kind == "fixed":
                e = ttk.Entry(cell, textvariable=var, width=22, state="readonly")
                e.pack(side="left", fill="x", expand=True)
            else:
                ttk.Entry(cell, textvariable=var, width=22).pack(
                    side="left", fill="x", expand=True)

    def _line_form(self, parent):
        form = ttk.Frame(parent)
        form.pack(fill="x")
        specs = [("F", 14), ("N", 10), ("O", 16), ("P", 12),
                 ("Q", 16), ("T", 12), ("W", 12), ("X", 12)]
        for i, (key, width) in enumerate(specs):
            cell = ttk.Frame(form)
            cell.grid(row=0, column=i, padx=4, sticky="w")
            label = HEADER_BY_KEY[key]
            if key in REQUIRED_KEYS:
                label = "* " + label
            ttk.Label(cell, text=label).pack(anchor="w")
            var = tk.StringVar()
            self.line_vars[key] = var
            row = ttk.Frame(cell)
            row.pack()
            ttk.Entry(row, textvariable=var, width=width).pack(side="left")
            if key == "Q":
                ttk.Button(row, text="찾기", width=5,
                           command=self._pick_fsc).pack(side="left", padx=2)
        # 단가 -> 금액 자동
        self.line_vars["W"].trace_add("write", lambda *_: self._auto_amount())

        btns = ttk.Frame(parent)
        btns.pack(fill="x", pady=(6, 6))
        ttk.Label(btns, text="납품요청일은 20261026 또는 2026-10-26 형식",
                  foreground="#777").pack(side="left")
        self.btn_add = ttk.Button(btns, text="행 추가", command=self._add_line)
        self.btn_add.pack(side="right")
        ttk.Button(btns, text="입력칸 비우기",
                   command=self._clear_line_form).pack(side="right", padx=6)

    def _line_table(self, parent):
        wrap = ttk.Frame(parent)
        wrap.pack(fill="both", expand=True)
        cols = ["No"] + LINE_KEYS
        self.tree = ttk.Treeview(wrap, columns=cols, show="headings", height=12)
        self.tree.heading("No", text="No")
        self.tree.column("No", width=40, anchor="center")
        widths = {"F": 110, "N": 80, "O": 130, "P": 100,
                  "Q": 130, "T": 100, "W": 110, "X": 110}
        for k in LINE_KEYS:
            self.tree.heading(k, text=HEADER_BY_KEY[k])
            self.tree.column(k, width=widths[k], anchor="w")
        vs = ttk.Scrollbar(wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vs.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vs.pack(side="left", fill="y")
        self.tree.bind("<Double-1>", lambda e: self._edit_line())

        tb = ttk.Frame(parent)
        tb.pack(fill="x", pady=(6, 0))
        ttk.Button(tb, text="선택 행 수정", command=self._edit_line).pack(side="left")
        ttk.Button(tb, text="선택 행 복제", command=self._dup_line).pack(side="left", padx=6)
        ttk.Button(tb, text="선택 행 삭제", command=self._del_line).pack(side="left")
        ttk.Button(tb, text="전체 삭제", command=self._clear_lines).pack(side="left", padx=6)
        self.line_count = ttk.Label(tb, text="0 행")
        self.line_count.pack(side="right")

    # ---------- 동작
    def _pick_master(self):
        p = filedialog.askopenfilename(
            title="통합양식 파일 선택",
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("모든 파일", "*.*")])
        if p:
            self.master_path.set(p)
            self._load_master()

    def _pick_partner(self, which):
        if not self.md:
            return
        if which == "sold":
            rows, title, key = self.md.sold_to, "판매처 선택", "B"
        else:
            rows, title, key = self.md.ship_to, "인도처 선택", "C"
        dlg = PickerDialog(self, title, ("코드", "이름", "주소"),
                           (90, 220, 220), rows)
        self.wait_window(dlg)
        if dlg.result:
            self.common_vars[key].set(dlg.result)

    def _show_partner_name(self, key):
        if not self.md:
            return
        code = self.common_vars[key].get().strip()
        rows = self.md.sold_to if key == "B" else self.md.ship_to
        name = next((r[1] for r in rows if r[0] == code), "")
        lbl = getattr(self, "_name_%s" % key, None)
        if lbl is not None:
            lbl.config(text=name if name else ("코드 없음" if code else ""),
                       foreground="#0a6" if name else "#c00")

    def _pick_fsc(self):
        if not self.md:
            return
        dlg = PickerDialog(self, "자재코드(FSC) 선택",
                           ("FSC", "VER", "모델명", "설명", "상태"),
                           (120, 45, 110, 300, 90), self.md.fsc)
        self.wait_window(dlg)
        if dlg.result:
            self.line_vars["Q"].set(dlg.result)

    def _auto_amount(self):
        w = parse_int(self.line_vars["W"].get())
        if w is not None:
            self.line_vars["X"].set(str(w))

    def _clear_line_form(self):
        for k in LINE_KEYS:
            self.line_vars[k].set("")
        self.editing_index = None
        self.btn_add.config(text="행 추가")

    def _validate_line(self, data):
        errs = []
        for k in LINE_KEYS:
            if k in REQUIRED_KEYS and not data[k].strip():
                errs.append("%s(%s) 은(는) 필수입니다." % (HEADER_BY_KEY[k], k))
        q = data["Q"].strip()
        if q and self.md and q not in self.md.fsc_codes:
            errs.append("자재코드 '%s' 은(는) FSC 목록에 없습니다." % q)
        if data["T"].strip() and parse_date(data["T"]) is None:
            errs.append("납품요청일 형식이 올바르지 않습니다. (예: 20261026)")
        for k in ("W", "X"):
            if data[k].strip() and parse_int(data[k]) is None:
                errs.append("%s 은(는) 숫자여야 합니다." % HEADER_BY_KEY[k])
        return errs

    def _add_line(self):
        data = {k: self.line_vars[k].get().strip() for k in LINE_KEYS}
        errs = self._validate_line(data)
        if errs:
            messagebox.showwarning("확인 필요", "\n".join(errs))
            return
        if self.editing_index is None:
            self.lines.append(data)
        else:
            self.lines[self.editing_index] = data
        self._refresh_tree()
        keep = {k: data[k] for k in ("F", "N", "O", "P", "T")}
        self._clear_line_form()
        for k, v in keep.items():          # 다음 행 입력 편의를 위해 유지
            self.line_vars[k].set(v)

    def _selected_index(self):
        sel = self.tree.selection()
        if not sel:
            return None
        return self.tree.index(sel[0])

    def _edit_line(self):
        i = self._selected_index()
        if i is None:
            return
        for k in LINE_KEYS:
            self.line_vars[k].set(self.lines[i][k])
        self.editing_index = i
        self.btn_add.config(text="수정 반영")

    def _dup_line(self):
        i = self._selected_index()
        if i is None:
            return
        self.lines.insert(i + 1, dict(self.lines[i]))
        self._refresh_tree()

    def _del_line(self):
        i = self._selected_index()
        if i is None:
            return
        del self.lines[i]
        self._clear_line_form()
        self._refresh_tree()

    def _clear_lines(self):
        if self.lines and messagebox.askyesno("확인", "품목 라인을 모두 지울까요?"):
            self.lines = []
            self._clear_line_form()
            self._refresh_tree()

    def _refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        for n, d in enumerate(self.lines, start=1):
            self.tree.insert("", "end", values=[n] + [d[k] for k in LINE_KEYS])
        self.line_count.config(text="%d 행" % len(self.lines))

    # ---------- 출력
    def _collect_common(self):
        out, errs = {}, []
        for k in COMMON_KEYS:
            raw = self.common_vars[k].get().strip()
            if k in ("A", "D", "H", "K", "Y"):
                raw = combo_code(raw)
            if k in REQUIRED_KEYS and not raw:
                errs.append("%s(%s) 은(는) 필수입니다." % (HEADER_BY_KEY[k], k))
            out[k] = raw
        if out.get("B") and self.md and out["B"] not in {r[0] for r in self.md.sold_to}:
            errs.append("판매처코드 '%s' 은(는) 목록에 없습니다." % out["B"])
        if out.get("C") and self.md and out["C"] not in {r[0] for r in self.md.ship_to}:
            errs.append("인도처코드 '%s' 은(는) 목록에 없습니다." % out["C"])
        for k in ("G", "J"):
            if out[k]:
                d = parse_date(out[k])
                if d is None:
                    errs.append("%s 형식이 올바르지 않습니다." % HEADER_BY_KEY[k])
                else:
                    out[k] = d.strftime("%Y%m%d")
        return out, errs

    def _build_rows(self, common):
        rows = []
        for d in self.lines:
            row = {}
            for k in COMMON_KEYS:
                v = common[k]
                if k in ("B", "C", "D", "E", "U"):
                    iv = parse_int(v)
                    v = iv if iv is not None else (v or None)
                elif k == "R":
                    v = 1                       # 오더수량은 항상 1
                row[k] = v if v != "" else None
            for k in LINE_KEYS:
                v = d[k]
                if k == "T":
                    v = parse_date(v)
                elif k in ("F", "W", "X"):
                    iv = parse_int(v)
                    v = iv if iv is not None else (v or None)
                row[k] = v if v != "" else None
            rows.append(row)
        return rows

    def _export(self):
        if not self.md:
            messagebox.showerror("오류", "먼저 양식 파일을 읽어주세요.")
            return
        if not self.lines:
            messagebox.showwarning("확인 필요", "품목 라인이 없습니다.")
            return
        common, errs = self._collect_common()
        for n, d in enumerate(self.lines, start=1):
            for e in self._validate_line(d):
                errs.append("%d행: %s" % (n, e))
        if errs:
            messagebox.showwarning("확인 필요", "\n".join(errs[:15]))
            return

        default = "CSP_주문접수_%s.xlsx" % dt.datetime.now().strftime("%Y%m%d_%H%M")
        out = filedialog.asksaveasfilename(
            title="저장 위치", defaultextension=".xlsx",
            initialfile=default, filetypes=[("Excel", "*.xlsx")])
        if not out:
            return
        try:
            build_output(self.master_path.get(), self._build_rows(common), out)
        except Exception as e:
            messagebox.showerror("오류", "파일 생성에 실패했습니다.\n\n%s" % e)
            return
        self.status.config(text="생성 완료 : %s" % out)
        messagebox.showinfo("완료", "%d행이 생성되었습니다.\n\n%s" % (len(self.lines), out))

    # ---------- 설정 저장 / 복원
    def _settings_path(self):
        return os.path.join(app_dir(), SETTINGS_NAME)

    def _save_settings(self, silent=False):
        data = {"master": self.master_path.get(),
                "common": {k: v.get() for k, v in self.common_vars.items()}}
        try:
            with open(self._settings_path(), "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            if not silent:
                self.status.config(text="공통값을 저장했습니다.")
        except Exception:
            pass

    def _load_settings(self):
        try:
            with open(self._settings_path(), encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            return
        for k, v in data.get("common", {}).items():
            if k in self.common_vars:
                self.common_vars[k].set(v)


if __name__ == "__main__":
    App().mainloop()
